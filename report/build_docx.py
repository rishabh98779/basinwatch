#!/usr/bin/env python3
"""Build the BasinWatch project report DOCX from the markdown drafts.

Two-pass build: pass 1 renders with placeholder TOC/LOF/LOT, we convert to PDF
and locate every heading and caption, then pass 2 renders identical content with
real page numbers. Because both passes contain the same number of TOC/LOF/LOT
lines, pagination is identical and the page numbers stay correct.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPORT = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(REPORT)
ASSETS = os.path.join(ROOT, "assets")
BANNER = os.path.join(ASSETS, "cu_banner.png")

RED = RGBColor(0xC0, 0x1B, 0x22)
MAROON = RGBColor(0x8B, 0x00, 0x00)

def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)

def para_bottom_border(p, color="C01B22", size=18):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), color)
    pbdr.append(bottom); pPr.append(pbdr)

def add_field(paragraph, instr):
    r = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = " "
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2); r._r.append(t); r._r.append(f3)

def add_page_number_footer(section):
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")
    for run in p.runs:
        run.font.name = 'Times New Roman'; run.font.size = Pt(11)

def add_header_banner(section):
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    if os.path.exists(BANNER):
        p.add_run().add_picture(BANNER, width=Inches(6.3))
    rule = header.add_paragraph()
    rule.paragraph_format.space_before = Pt(0); rule.paragraph_format.space_after = Pt(2)
    para_bottom_border(rule, size=20)

def build_styles(doc):
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    pf = st.paragraph_format
    pf.line_spacing = 1.4; pf.space_after = Pt(8); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color, before, after in [
        ('Heading 1', 17, MAROON, 16, 10),
        ('Heading 2', 14, MAROON, 12, 6),
        ('Heading 3', 12.5, RGBColor(0x33,0x33,0x33), 10, 4)]:
        s = doc.styles[name]
        s.font.name = 'Times New Roman'; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    cap = doc.styles.add_style('FigCaption', 1)
    cap.font.name = 'Times New Roman'; cap.font.size = Pt(10); cap.font.italic = True
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2); cap.paragraph_format.space_after = Pt(12)
    tc = doc.styles.add_style('TabCaption', 1)
    tc.font.name = 'Times New Roman'; tc.font.size = Pt(10); tc.font.italic = True
    tc.paragraph_format.space_before = Pt(2); tc.paragraph_format.space_after = Pt(12)
    code = doc.styles.add_style('CodeBlock', 1)
    code.font.name = 'Consolas'; code.font.size = Pt(9)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0); code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tst = doc.styles.add_style('TocEntry', 1)
    tst.font.name = 'Times New Roman'; tst.font.size = Pt(11)
    tst.paragraph_format.space_after = Pt(2); tst.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_inline(paragraph, text):
    for tok in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(10.5)
        else:
            paragraph.add_run(tok)

def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h); r.bold = True; r.font.size = Pt(10.5); r.font.name = 'Times New Roman'
        set_cell_shading(hdr[i], 'E8E8E8')
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if i >= len(cells): break
            cells[i].text = ''
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_inline(p, val)
            for r in p.runs:
                r.font.size = Pt(10.5)
                if not r.font.name: r.font.name = 'Times New Roman'
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.keep_with_next = True
    return t

def add_figure(doc, path, caption):
    if not os.path.exists(path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"[IMAGE MISSING: {path}]").font.color.rgb = RED
        doc.add_paragraph(caption, style='FigCaption')
        return
    from PIL import Image
    w, h = Image.open(path).size
    width = min(6.0, w / 120.0)
    height = width * h / w
    MAX_H = 6.9
    if height > MAX_H:
        height = MAX_H
        width = height * w / h
    # image and caption share one paragraph so they can never separate
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width))
    p.add_run().add_break()
    r = p.add_run(caption); r.italic = True; r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(12)

def add_title_page(doc, lines):
    for ln in lines:
        ln = ln.strip()
        if not ln: continue
        if ln == '[red rule]':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para_bottom_border(p, color="C01B22", size=24)
            p.paragraph_format.space_after = Pt(10); continue
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        if ln == 'CHANDIGARH UNIVERSITY':
            r = p.add_run(ln); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = MAROON
        elif ln == 'PROJECT REPORT':
            r = p.add_run(ln); r.bold = True; r.font.size = Pt(18)
        elif ln.startswith('"'):
            r = p.add_run(ln); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RED
        else:
            add_inline(p, ln)
            for r in p.runs: r.font.size = Pt(12)
    doc.add_page_break()

def toc_line(doc, text, page, indent=0, bold=False):
    p = doc.add_paragraph(style='TocEntry')
    p.paragraph_format.left_indent = Inches(0.15 + indent * 0.28)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.15), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(text); r.bold = bold
    if bold: r.font.size = Pt(11.5)
    p.add_run("\t"); p.add_run(str(page))

def scan_entries(stream):
    entries = []; fence = False; h1 = h2 = fig = tab = 0
    for line in stream.split('\n'):
        s = line.strip()
        if s.startswith('```'):
            fence = not fence; continue
        if fence: continue
        if s.startswith('### '): continue
        if s.startswith('## '):
            h2 += 1; entries.append({'kind': 'H2', 'text': s[3:], 'key': f'H2:{h2}'}); continue
        if s.startswith('# '):
            h1 += 1; entries.append({'kind': 'H1', 'text': s[2:], 'key': f'H1:{h1}'}); continue
        m = re.match(r'\[\[FIG:.+?\|(.+?)\]\]', s)
        if m:
            fig += 1; entries.append({'kind': 'FIG', 'text': m.group(1), 'key': f'FIG:{fig}'}); continue
        if re.match(r'\*Figure [0-9]', s) and s.endswith('*'):
            fig += 1; entries.append({'kind': 'FIG', 'text': s[1:-1], 'key': f'FIG:{fig}'}); continue
        if re.match(r'\*Table [0-9]', s) and s.endswith('*'):
            tab += 1; entries.append({'kind': 'TABLE', 'text': s[1:-1], 'key': f'TABLE:{tab}'}); continue
    return entries

def parse_markdown(doc, text, pages, is_pass2):
    lines = text.split('\n'); i = 0
    while i < len(lines):
        line = lines[i]; s = line.strip()
        if s.startswith('%%%') and s.endswith('%%%'):
            if len(doc.paragraphs) > 0: doc.add_page_break()
            i += 1; continue
        if s == '[PAGEBREAK]':
            doc.add_page_break(); i += 1; continue
        if s == '[red rule]':
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
            para_bottom_border(p, color="C01B22", size=20)
            i += 1; continue
        if s == '[[TOC]]':
            for e in TOC:
                toc_line(doc, e['text'], pages.get(e['key'], '00') if is_pass2 else '00',
                         indent=0 if e['kind'] == 'H1' else 1, bold=(e['kind'] == 'H1'))
            i += 1; continue
        if s == '[[LOF]]':
            for e in LOF:
                toc_line(doc, e['text'], pages.get(e['key'], '00') if is_pass2 else '00')
            i += 1; continue
        if s == '[[LOT]]':
            for e in LOT:
                toc_line(doc, e['text'], pages.get(e['key'], '00') if is_pass2 else '00')
            i += 1; continue
        if s == '[TITLE_PAGE]':
            i += 1; block = []
            while lines[i].strip() != '[END_TITLE_PAGE]':
                block.append(lines[i]); i += 1
            add_title_page(doc, block); i += 1; continue
        if s.startswith('[[FIG:'):
            m = re.match(r'\[\[FIG:(.+?)\|(.+?)\]\]', s)
            path, cap = m.group(1), m.group(2)
            if not os.path.isabs(path):
                cand = os.path.join(ROOT, path)
                if not os.path.exists(cand):
                    cand = os.path.join(REPORT, path)
                if not os.path.exists(cand):
                    cand = os.path.join(REPORT, 'diagrams', os.path.basename(path))
                if not os.path.exists(cand):
                    cand = os.path.join(ROOT, 'assets', os.path.basename(path))
                path = cand
            add_figure(doc, path, cap); i += 1; continue
        if s.startswith('```'):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            for cl in buf:
                doc.add_paragraph(cl if cl.strip() else '', style='CodeBlock')
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue
        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i+1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not re.match(r'^[\s\-:]+$', ''.join(cells)): rows.append(cells)
                i += 1
            if rows: add_table(doc, rows)
            continue
        if s.startswith('### '): doc.add_heading(s[4:], level=3); i += 1; continue
        if s.startswith('## '): doc.add_heading(s[3:], level=2); i += 1; continue
        if s.startswith('# '): doc.add_heading(s[2:], level=1); i += 1; continue
        if s.startswith('*') and s.endswith('*') and len(s) > 2 and not s.startswith('**'):
            if re.match(r'\*Figure', s): doc.add_paragraph(s[1:-1], style='FigCaption')
            elif re.match(r'\*Table', s): doc.add_paragraph(s[1:-1], style='TabCaption')
            else:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
                p.add_run(s[1:-1]).italic = True
            i += 1; continue
        if s == '---': i += 1; continue
        if s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, s[2:]); i += 1; continue
        if not s: i += 1; continue
        p = doc.add_paragraph(); add_inline(p, s); i += 1

TOC = []; LOF = []; LOT = []

def assemble_stream():
    d00 = open(f"{REPORT}/draft_00_frontmatter_pages.md", encoding='utf-8').read()
    d01 = open(f"{REPORT}/draft_01_frontmatter_ch1_ch2.md", encoding='utf-8').read()
    d02 = open(f"{REPORT}/draft_02_ch3.md", encoding='utf-8').read()
    d03 = open(f"{REPORT}/draft_03_ch4.md", encoding='utf-8').read()
    d04 = open(f"{REPORT}/draft_04_ch5_ch6.md", encoding='utf-8').read()
    d05 = open(f"{REPORT}/draft_05_ch7_ch8_ch9.md", encoding='utf-8').read()
    d06 = open(f"{REPORT}/draft_06_ch10_ch11_ch12_refs_appendix.md", encoding='utf-8').read()
    head00, abbrev = d00.split('%%% ABBREVIATIONS %%%', 1)
    abbrev = '%%% ABBREVIATIONS %%%' + abbrev
    for tok in ('[[TITLEPAGE]]', '[[CERTIFICATE]]', '[[DECLARATION]]'):
        d01 = d01.replace(tok, '')
    front, chapters = d01.split('%%% CHAPTER 1 %%%', 1)
    toc_block = ("%%% TABLE OF CONTENTS %%%\n# TABLE OF CONTENTS\n[[TOC]]\n"
                 "%%% LIST OF FIGURES %%%\n# LIST OF FIGURES\n[[LOF]]\n"
                 "%%% LIST OF TABLES %%%\n# LIST OF TABLES\n[[LOT]]\n")
    return (head00 + front + abbrev + toc_block +
            '%%% CHAPTER 1 %%%' + chapters + d02 + d03 + d04 + d05 + d06)

def build(out, pages, is_pass2):
    doc = Document(); build_styles(doc)
    sec = doc.sections[0]
    sec.page_height = Inches(11.69); sec.page_width = Inches(8.27)
    sec.left_margin = Inches(1.1); sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8)
    sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.4)
    sec.different_first_page_header_footer = True
    add_header_banner(sec); add_page_number_footer(sec)
    parse_markdown(doc, assemble_stream(), pages, is_pass2)
    doc.save(out)

def main():
    global TOC, LOF, LOT
    stream = assemble_stream()
    entries = scan_entries(stream)
    TOC = [e for e in entries if e['kind'] in ('H1', 'H2')]
    LOF = [e for e in entries if e['kind'] == 'FIG']
    LOT = [e for e in entries if e['kind'] == 'TABLE']
    print(f"TOC={len(TOC)} LOF={len(LOF)} LOT={len(LOT)}")
    build(f"{REPORT}/_pass1.docx", {}, False)
    os.system("cd %s && LD_LIBRARY_PATH=/usr/lib/libreoffice/program /usr/bin/soffice --headless "
              "--convert-to pdf _pass1.docx --outdir . >/dev/null 2>&1" % REPORT)
    import pymupdf
    doc = pymupdf.open(f"{REPORT}/_pass1.pdf")
    pages = {}; cursor = 0
    listing = set()
    for pi in range(len(doc)):
        if doc[pi].get_text().count('....') > 5:
            listing.add(pi)
    def locate(needle, start):
        for pi in range(start, len(doc)):
            if pi in listing: continue
            if needle and needle in re.sub(r'\s+', ' ', doc[pi].get_text()):
                return pi + 1
        return None
    def heading_page(text):
        for pi in range(len(doc)):
            lines = [l.strip() for l in doc[pi].get_text().split('\n') if l.strip()]
            if lines and lines[0] == text:
                return pi + 1
        return None
    for e in entries:
        if e['text'] in ('TABLE OF CONTENTS', 'LIST OF FIGURES', 'LIST OF TABLES'):
            pages[e['key']] = heading_page(e['text']); continue
        needle = re.sub(r'\s+', ' ', e['text']).strip()
        found = locate(needle, cursor)
        if found is None:
            found = locate(needle, 0)
        if found is not None:
            cursor = found - 1; pages[e['key']] = found
    print("located", len(pages), "of", len(entries))
    build(f"{REPORT}/BasinWatch_Project_Report.docx", pages, True)
    print("saved final docx")

if __name__ == '__main__':
    main()
